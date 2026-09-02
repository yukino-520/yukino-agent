from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path
from typing import Any


# 作用：表示工作区路径、文件格式或资源大小违反安全限制。
# 参数：无。
class WorkspaceError(ValueError):
    pass


# 作用：在固定字符预算内逐段汇总文档文本。
# 参数：无。
class _TextAccumulator:
    # 作用：初始化文本上限、片段列表和当前长度。
    # 参数 limit：本次读取、搜索或缓冲允许返回的最大数量。
    def __init__(self, limit: int) -> None:
        self.limit = limit
        self.parts: list[str] = []
        self.size = 0

    # 作用：追加可容纳的文本片段，并报告内容是否完整写入。
    # 参数 value：待解析、清洗、转义或递归处理的输入值。
    # 参数 separator：已有文本和新片段之间使用的分隔符。
    def add(self, value: object, *, separator: str = "\n") -> bool:
        if self.size >= self.limit:
            return False
        text = str(value)
        if self.parts:
            remaining = self.limit - self.size
            prefix = separator[:remaining]
            self.parts.append(prefix)
            self.size += len(prefix)
        remaining = self.limit - self.size
        if remaining <= 0:
            return False
        piece = text[:remaining]
        self.parts.append(piece)
        self.size += len(piece)
        return len(text) <= remaining

    # 作用：拼接并返回累计的受限长度文本。
    # 参数：无。
    def build(self) -> str:
        return "".join(self.parts)


# 作用：在专属根目录中提供防越界文件操作和受限文档解析。
# 参数：无。
class SafeWorkspace:
    MAX_READ_BYTES = 2 * 1024 * 1024
    MAX_WRITE_BYTES = 2 * 1024 * 1024
    MAX_DOCUMENT_BYTES = 16 * 1024 * 1024
    MAX_DOCUMENT_OUTPUT_CHARS = 200_000
    MAX_ARCHIVE_ENTRIES = 5_000
    MAX_ARCHIVE_MEMBER_BYTES = 16 * 1024 * 1024
    MAX_ARCHIVE_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
    MAX_PDF_PAGES = 1_000

    # 作用：解析并创建工作区根目录。
    # 参数 root：安全工作区限定的根目录。
    def __init__(self, root: str | Path) -> None:
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    # 作用：解析相对路径并阻止目录穿越。
    # 参数 relative：相对于安全工作区根目录的路径。
    # 参数 must_exist：解析路径时是否要求目标已经存在。
    def resolve(self, relative: str = ".", *, must_exist: bool = False) -> Path:
        candidate = (self.root / relative.strip().lstrip("/")).resolve()
        if not candidate.is_relative_to(self.root):
            raise WorkspaceError("路径超出 AGI Yukino 工作区。")
        if must_exist and not candidate.exists():
            raise WorkspaceError("文件不存在。")
        return candidate

    # 作用：列出工作区目录中有限数量的文件和子目录。
    # 参数 relative：相对于安全工作区根目录的路径。
    def list(self, relative: str = ".") -> list[dict[str, Any]]:
        directory = self.resolve(relative, must_exist=True)
        if not directory.is_dir():
            raise WorkspaceError("目标不是目录。")
        return [
            {
                "name": item.name,
                "path": str(item.relative_to(self.root)),
                "kind": "directory" if item.is_dir() else "file",
                "size": item.stat().st_size if item.is_file() else 0,
            }
            for item in sorted(directory.iterdir(), key=lambda value: (not value.is_dir(), value.name.lower()))[:500]
        ]

    # 作用：读取大小受限的 UTF-8 文本文件。
    # 参数 relative：相对于安全工作区根目录的路径。
    def read(self, relative: str) -> dict[str, Any]:
        source = self.resolve(relative, must_exist=True)
        if not source.is_file():
            raise WorkspaceError("目标不是文件。")
        if source.stat().st_size > self.MAX_READ_BYTES:
            raise WorkspaceError("文件超过 2 MiB 读取限制。")
        raw = source.read_bytes()
        try:
            content = raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise WorkspaceError("这个接口只读取 UTF-8 文本；二进制文件请使用媒体接口。") from exc
        return {"path": str(source.relative_to(self.root)), "content": content, "size": len(raw)}

    # 作用：在工作区内写入受限大小内容，覆盖时要求显式许可。
    # 参数 relative：相对于安全工作区根目录的路径。
    # 参数 content：待写入文件的文本内容。
    # 参数 overwrite：目标已存在时是否允许覆盖写入。
    def write(self, relative: str, content: str, *, overwrite: bool = False) -> dict[str, Any]:
        encoded = content.encode("utf-8")
        if len(encoded) > self.MAX_WRITE_BYTES:
            raise WorkspaceError("写入内容超过 2 MiB 限制。")
        target = self.resolve(relative)
        if target.exists() and not overwrite:
            raise WorkspaceError("文件已存在；覆盖需要明确确认。")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(encoded)
        return {"path": str(target.relative_to(self.root)), "size": len(encoded), "overwritten": overwrite}

    # 作用：递归检索工作区文本文件并限制命中数量。
    # 参数 query：能力、文件、网页或图谱检索使用的查询文本。
    # 参数 relative：相对于安全工作区根目录的路径。
    def search(self, query: str, relative: str = ".") -> list[dict[str, Any]]:
        if not query.strip():
            raise WorkspaceError("搜索词不能为空。")
        root = self.resolve(relative, must_exist=True)
        pattern = re.compile(re.escape(query), re.IGNORECASE)
        results: list[dict[str, Any]] = []
        paths = [root] if root.is_file() else root.rglob("*")
        for source in paths:
            if not source.is_file() or source.stat().st_size > self.MAX_READ_BYTES:
                continue
            try:
                lines = source.read_text(encoding="utf-8").splitlines()
            except (OSError, UnicodeDecodeError):
                continue
            for line_number, line in enumerate(lines, 1):
                if pattern.search(line):
                    results.append(
                        {"path": str(source.relative_to(self.root)), "line": line_number, "text": line[:300]}
                    )
                    if len(results) >= 200:
                        return results
        return results

    # 作用：按扩展名安全提取文本、PDF 和 Office 文档内容。
    # 参数 relative：相对于安全工作区根目录的路径。
    def read_document(self, relative: str) -> dict[str, Any]:
        source = self.resolve(relative, must_exist=True)
        self._validate_document_source(source)
        suffix = source.suffix.lower()
        if suffix in {".txt", ".md", ".py", ".html", ".css", ".js", ".yaml", ".yml"}:
            return self.read(relative)
        if suffix == ".json":
            try:
                value = json.loads(source.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise WorkspaceError("JSON 文档无法解析。") from exc
            return {
                "path": relative,
                "content": json.dumps(value, ensure_ascii=False, indent=2)[
                    : self.MAX_DOCUMENT_OUTPUT_CHARS
                ],
            }
        if suffix == ".csv":
            output = _TextAccumulator(self.MAX_DOCUMENT_OUTPUT_CHARS)
            try:
                with source.open(encoding="utf-8", newline="") as handle:
                    for index, row in enumerate(csv.reader(handle)):
                        if index >= 1_000 or not output.add(" | ".join(row[:200])):
                            break
            except (OSError, UnicodeDecodeError, csv.Error) as exc:
                raise WorkspaceError("CSV 文档无法解析。") from exc
            return {"path": relative, "content": output.build()}
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise WorkspaceError("读取 PDF 需要安装 pypdf。") from exc
            try:
                reader = PdfReader(str(source))
                if reader.is_encrypted:
                    raise WorkspaceError("不读取加密 PDF。")
                page_count = len(reader.pages)
                if page_count > self.MAX_PDF_PAGES:
                    raise WorkspaceError("PDF 页数超过 1000 页安全限制。")
                output = _TextAccumulator(self.MAX_DOCUMENT_OUTPUT_CHARS)
                for page in reader.pages[:100]:
                    if not output.add(page.extract_text() or "", separator="\n\n"):
                        break
            except WorkspaceError:
                raise
            except Exception as exc:
                raise WorkspaceError("PDF 文档无法安全解析。") from exc
            return {"path": relative, "content": output.build(), "pages": page_count}
        if suffix == ".docx":
            self._preflight_office_archive(source)
            try:
                from docx import Document
            except ImportError as exc:
                raise WorkspaceError("读取 DOCX 需要安装 python-docx。") from exc
            try:
                document = Document(str(source))
                output = _TextAccumulator(self.MAX_DOCUMENT_OUTPUT_CHARS)
                for paragraph in document.paragraphs:
                    if not output.add(paragraph.text):
                        break
            except Exception as exc:
                raise WorkspaceError("DOCX 文档无法安全解析。") from exc
            return {"path": relative, "content": output.build()}
        if suffix in {".xlsx", ".xls"}:
            if suffix == ".xls":
                try:
                    import xlrd
                except ImportError as exc:
                    raise WorkspaceError("读取 XLS 需要安装 xlrd。") from exc
                try:
                    workbook = xlrd.open_workbook(source)
                    output = _TextAccumulator(self.MAX_DOCUMENT_OUTPUT_CHARS)
                    for sheet in workbook.sheets()[:30]:
                        if not output.add(f"## {sheet.name}"):
                            break
                        for index in range(min(sheet.nrows, 1_000)):
                            if not output.add(
                                " | ".join(
                                    str(value) for value in sheet.row_values(index)[:200]
                                )
                            ):
                                break
                except Exception as exc:
                    raise WorkspaceError("XLS 文档无法安全解析。") from exc
                return {
                    "path": relative,
                    "content": output.build(),
                    "sheets": workbook.nsheets,
                }
            self._preflight_office_archive(source)
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise WorkspaceError("读取 XLSX 需要安装 openpyxl。") from exc
            workbook = None
            try:
                workbook = load_workbook(source, read_only=True, data_only=True)
                output = _TextAccumulator(self.MAX_DOCUMENT_OUTPUT_CHARS)
                for sheet in workbook.worksheets[:30]:
                    if not output.add(f"## {sheet.title}"):
                        break
                    for row in sheet.iter_rows(
                        max_row=min(sheet.max_row, 1_000),
                        max_col=min(sheet.max_column, 200),
                        values_only=True,
                    ):
                        if not output.add(
                            " | ".join(
                                "" if value is None else str(value) for value in row
                            )
                        ):
                            break
                sheet_count = len(workbook.sheetnames)
            except Exception as exc:
                raise WorkspaceError("XLSX 文档无法安全解析。") from exc
            finally:
                if workbook is not None:
                    workbook.close()
            return {
                "path": relative,
                "content": output.build(),
                "sheets": sheet_count,
            }
        if suffix == ".pptx":
            self._preflight_office_archive(source)
            try:
                from pptx import Presentation
            except ImportError as exc:
                raise WorkspaceError("读取 PPTX 需要安装 python-pptx。") from exc
            try:
                presentation = Presentation(str(source))
                slide_count = len(presentation.slides)
                if slide_count > 500:
                    raise WorkspaceError("PPTX 页数超过 500 页安全限制。")
                output = _TextAccumulator(self.MAX_DOCUMENT_OUTPUT_CHARS)
                for index, slide in enumerate(presentation.slides, 1):
                    if index > 200 or not output.add(f"## Slide {index}"):
                        break
                    for shape in slide.shapes:
                        text = getattr(shape, "text", "")
                        if text and not output.add(text):
                            break
            except WorkspaceError:
                raise
            except Exception as exc:
                raise WorkspaceError("PPTX 文档无法安全解析。") from exc
            return {
                "path": relative,
                "content": output.build(),
                "slides": slide_count,
            }
        raise WorkspaceError(f"暂不支持 {suffix or '无扩展名'} 文档。")

    # 作用：确认文档是普通文件且未超过总大小限制。
    # 参数 source：待读取、迁移、验证或解析的源数据。
    def _validate_document_source(self, source: Path) -> None:
        if not source.is_file():
            raise WorkspaceError("目标不是文件。")
        try:
            size = source.stat().st_size
        except OSError as exc:
            raise WorkspaceError("无法读取文档信息。") from exc
        if size > self.MAX_DOCUMENT_BYTES:
            raise WorkspaceError("文档超过 16 MiB 解析限制。")

    # 作用：在解析 Office 文件前检查压缩条目，防止路径穿越和解压炸弹。
    # 参数 source：待读取、迁移、验证或解析的源数据。
    def _preflight_office_archive(self, source: Path) -> None:
        try:
            with zipfile.ZipFile(source) as archive:
                members = archive.infolist()
                if len(members) > self.MAX_ARCHIVE_ENTRIES:
                    raise WorkspaceError("Office 文档内部文件数量超过安全限制。")
                total_size = 0
                names: set[str] = set()
                for member in members:
                    normalized = member.filename.replace("\\", "/")
                    parts = normalized.split("/")
                    if (
                        not normalized
                        or normalized.startswith("/")
                        or ".." in parts
                        or normalized in names
                        or member.flag_bits & 0x1
                    ):
                        raise WorkspaceError("Office 文档包含不安全的归档条目。")
                    names.add(normalized)
                    if member.file_size > self.MAX_ARCHIVE_MEMBER_BYTES:
                        raise WorkspaceError("Office 文档内部单个文件超过安全限制。")
                    total_size += member.file_size
                    if total_size > self.MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                        raise WorkspaceError("Office 文档解压后超过 64 MiB 安全限制。")
        except WorkspaceError:
            raise
        except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
            raise WorkspaceError("Office 文档归档结构无效。") from exc
